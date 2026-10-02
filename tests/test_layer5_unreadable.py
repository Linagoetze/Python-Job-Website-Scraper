"""Layer 5's third outcome: the page was fetched and held no posting (SP4c).

A 200 response is not a reading. A client-rendered detail page arrives as a
title and a notice, and before SP4c Layer 5 read that as "checked, nothing
required" and, for the two fail-closed deferred states, as "checked, found
lacking": 24 stored rows were rejected for good against a shell.

The pages here are real in shape (a shell, a posting) and invented in content.
"""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Any

import pytest
import yaml
from openpyxl import load_workbook

from job_scraper import pipeline as pipeline_mod
from job_scraper.experience_filter import (
    EXPERIENCE_UNREADABLE,
    MIN_READABLE_CHARS,
    PAGE_STATE_KEY,
    PAGE_UNREADABLE,
    UNVERIFIED_KEY,
    apply_detail_filter,
    is_unreadable,
)
from job_scraper.filtering import _HYBRID_PENDING_REASON, _UNRESOLVED_PENDING_REASON
from job_scraper.pipeline import UnreadablePages, count_unreadable_pages, run_pipeline
from job_scraper.run import format_summary
from job_scraper.storage.db import JobStore
from job_scraper.storage.xlsx_store import write_xlsx
from tests.pages import posting
from tests.test_run_summary import _summary

_SHELL = (
    "<html><head><title>Analyst @ Contoso</title></head><body>"
    "<noscript>You need to enable JavaScript to run this app.</noscript>"
    '<div id="root"></div></body></html>'
)

# --- the decision ------------------------------------------------------------


def test_a_short_page_is_unreadable_at_the_threshold() -> None:
    assert is_unreadable("x" * (MIN_READABLE_CHARS - 1))
    assert not is_unreadable("x" * MIN_READABLE_CHARS)


def test_the_threshold_sits_in_the_gap_the_store_showed() -> None:
    """Stored shells were all 159 characters or fewer; readable postings 2,060 or more."""
    assert 159 < MIN_READABLE_CHARS < 2_060


def test_a_js_marker_makes_a_page_below_the_ceiling_unreadable() -> None:
    padded = "Please enable JavaScript to view this page. " + "Cookie notice. " * 80
    assert MIN_READABLE_CHARS < len(padded) < 2_000
    assert is_unreadable(padded)


def test_a_readable_posting_with_a_noscript_footer_stays_readable() -> None:
    page = "Responsibilities and requirements. " * 120 + " Please enable JavaScript."
    assert len(page) > 2_000
    assert not is_unreadable(page)


def test_a_long_page_without_a_marker_is_readable_whatever_it_says() -> None:
    """The tetrapak case (one cookie page, four postings) is not caught here, on purpose."""
    assert not is_unreadable("Accept all cookies. " * 200)


# --- Layer 5 ---------------------------------------------------------------


def _job(**over: Any) -> dict[str, Any]:
    return {
        "source_name": "contoso",
        "title": "Analyst",
        "location": "Berlin",
        "detail_url": "https://contoso.example/jobs/1",
        "apply_url": "",
        "matched_reasons": [],
        **over,
    }


def _detail(job: dict[str, Any], page: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    return apply_detail_filter([job], lambda _url: page)


def test_a_kept_unreadable_job_is_marked_and_carries_no_description() -> None:
    (kept,), excluded = _detail(_job(), _SHELL)
    assert not excluded
    assert kept["experience_level"] == EXPERIENCE_UNREADABLE
    assert kept["description_text"] == "", "a shell is not a description; it is also not scored"
    assert kept["description_fetched_at"] == ""
    assert kept[PAGE_STATE_KEY] == PAGE_UNREADABLE
    assert UNVERIFIED_KEY not in kept


def test_nothing_is_read_off_a_shell() -> None:
    """Fail open, as the owner decided: there is nothing to fail on."""
    page = "<p>8+ years of experience. A PhD is required.</p>"
    (kept,), excluded = _detail(_job(), page)
    assert kept["experience_level"] == EXPERIENCE_UNREADABLE
    assert not excluded


@pytest.mark.parametrize(
    ("location", "pending"),
    [
        ("2 Locations", _UNRESOLVED_PENDING_REASON),
        ("Faraway", _HYBRID_PENDING_REASON),
    ],
)
def test_a_deferred_job_against_a_shell_is_unverified_with_no_description(
    location: str, pending: str
) -> None:
    kept, (dropped,) = _detail(_job(location=location, matched_reasons=[pending]), _SHELL)
    assert not kept
    assert dropped[UNVERIFIED_KEY] is True
    assert dropped["description_text"] == ""


def _boom(_url: str) -> str:
    raise RuntimeError("network down")


def test_a_failed_fetch_is_marked_unchecked_not_unspecified() -> None:
    (kept,), excluded = apply_detail_filter([_job()], _boom)
    assert not excluded
    assert kept["experience_level"] == EXPERIENCE_UNREADABLE


def test_a_job_with_no_url_is_marked_unchecked_not_unspecified() -> None:
    (kept,), _ = _detail(_job(detail_url=""), _SHELL)
    assert kept["experience_level"] == EXPERIENCE_UNREADABLE


def test_unspecified_is_only_for_a_posting_read_in_full() -> None:
    (kept,), _ = _detail(_job(), posting("A role that states no requirement."))
    assert kept["experience_level"] == "unspecified"


def test_a_readable_page_is_still_read() -> None:
    (kept,), _ = _detail(
        _job(), posting("Requires 1 year of experience... 2+ years of experience.")
    )
    assert kept["experience_level"] == "junior (<=2yr)"


# --- the count behind the summary block -------------------------------------


def test_unreadable_pages_are_counted_per_source_against_pages_that_came_back() -> None:
    def j(source: str, state: str | None) -> dict[str, Any]:
        return {"source_name": source, **({PAGE_STATE_KEY: state} if state else {})}

    jobs = [
        j("undp", "unreadable"),
        j("undp", "unreadable"),
        j("undp", "read"),
        j("undp", "failed"),  # a request that did not complete is not a page
        j("undp", None),  # no URL
        j("fine", "read"),
        j("kognity", "unreadable"),
    ]
    assert count_unreadable_pages(jobs) == (
        UnreadablePages("kognity", 1, 1),
        UnreadablePages("undp", 2, 3),
    )


def test_the_block_renders_beside_the_other_warnings() -> None:
    text = format_summary(
        _summary(
            unreadable_pages=(UnreadablePages("undp", 16, 20), UnreadablePages("kognity", 1, 4))
        )
    )
    assert (
        "!  Unreadable pages: 2 sources had detail pages with no posting in them\n"
        "!  undp: 16 of 20 detail pages unreadable — those jobs are not experience-checked\n"
        "!  kognity: 1 of 4 detail pages unreadable — those jobs are not experience-checked"
    ) in text


def test_no_block_when_every_page_was_read() -> None:
    assert "Unreadable pages" not in format_summary(_summary())


# --- the sheet --------------------------------------------------------------


def test_the_sheet_shows_a_job_is_unchecked(tmp_path: Path) -> None:
    db, xlsx = tmp_path / "jobs.sqlite3", tmp_path / "jobs.xlsx"
    with JobStore(db) as store:
        run_id = store.begin_run()
        store.upsert_jobs(
            [
                {
                    "dedupe_key": "https://x/1",
                    "source_name": "acme",
                    "title": "Analyst",
                    "detail_url": "https://x/1",
                    "experience_level": EXPERIENCE_UNREADABLE,
                }
            ],
            run_id,
        )
        store.finish_run(run_id)
    write_xlsx(db, xlsx)
    ws = load_workbook(xlsx)["Jobs"]
    headers = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=headers.index("experience_level") + 1).value == (
        EXPERIENCE_UNREADABLE
    )


# --- the pipeline: re-checking across runs ----------------------------------

_SOURCE = "acme"
_LISTING = "https://acme.example/jobs"
_URL = f"{_LISTING}/a"


class _Site:
    """A detail page that can be switched between a shell, a posting and an error."""

    def __init__(self) -> None:
        self.page: str | Exception = _SHELL
        self.fetches: Counter[str] = Counter()

    def __call__(self, url: str, *a: Any, **k: Any) -> str:
        self.fetches[url] += 1
        if isinstance(self.page, Exception):
            raise self.page
        return self.page


@pytest.fixture
def env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    (tmp_path / "sources.yaml").write_text(
        yaml.dump({"sources": [{"name": _SOURCE, "url": _LISTING, "strategy": "static"}]}),
        encoding="utf-8",
    )
    (tmp_path / "rules.json").write_text(
        json.dumps({"locations": ["Berlin"], "non_place_locations": ["Worldwide"]}),
        encoding="utf-8",
    )
    job = {
        "source_name": _SOURCE,
        "title": "Data Analyst",
        "company": "Acme",
        "location": "Berlin",
        "listing_url": _LISTING,
        "detail_url": _URL,
        "apply_url": "",
        "raw_snippet": "Data Analyst Berlin",
    }
    site = _Site()
    extracted = [job]
    monkeypatch.setattr(
        pipeline_mod, "get_extractor", lambda name: lambda url, fetch_fn: list(extracted)
    )
    monkeypatch.setattr(pipeline_mod, "fetch_text", site)
    monkeypatch.setattr(pipeline_mod, "fetch_rendered", site)
    return {"tmp_path": tmp_path, "site": site, "extracted": extracted}


def _run(tmp_path: Path):
    return run_pipeline(
        sources_path=tmp_path / "sources.yaml",
        rules_path=tmp_path / "rules.json",
        out_db_path=tmp_path / "jobs.sqlite3",
        cache_path=tmp_path / "http_cache.sqlite3",
        delist_after=1,
        check_robots=False,
    )


def _stored(tmp_path: Path) -> dict[str, dict[str, Any]]:
    with JobStore(tmp_path / "jobs.sqlite3") as store:
        return {j["dedupe_key"]: j for j in store.all_jobs()}


def test_an_unreadable_job_is_stored_unchecked_and_reported(env: dict[str, Any]) -> None:
    summary = _run(env["tmp_path"])
    job = _stored(env["tmp_path"])[_URL]
    assert (job["status"], job["experience_level"], job["description_text"]) == (
        "new",
        EXPERIENCE_UNREADABLE,
        "",
    )
    assert summary.unreadable_pages == (UnreadablePages(_SOURCE, 1, 1),)


def test_an_unreadable_job_is_fetched_again_until_it_reads(env: dict[str, Any]) -> None:
    """Pins the request count, as test_logging_costs_no_extra_http_request does.

    Before SP4c a stored job was never detail-fetched again, so a shell was final.
    """
    site, tmp_path = env["site"], env["tmp_path"]
    _run(tmp_path)
    assert site.fetches[_URL] == 1

    second = _run(tmp_path)
    assert site.fetches[_URL] == 2, "a stored job with no description is re-fetched"
    assert second.jobs_stored_rechecked == 1
    assert second.jobs_new_checked == 0
    assert _stored(tmp_path)[_URL]["experience_level"] == EXPERIENCE_UNREADABLE

    site.page = posting("A great opportunity. No requirement is stated.")
    third = _run(tmp_path)
    assert site.fetches[_URL] == 3
    assert third.unreadable_pages == ()
    job = _stored(tmp_path)[_URL]
    assert job["experience_level"] == "unspecified"
    assert job["description_text"].startswith("A great opportunity")

    _run(tmp_path)
    assert site.fetches[_URL] == 3, "once read, it is skipped again"


def test_a_failed_recheck_does_not_erase_the_unchecked_mark(env: dict[str, Any]) -> None:
    site, tmp_path = env["site"], env["tmp_path"]
    _run(tmp_path)
    site.page = RuntimeError("network down")
    _run(tmp_path)
    assert _stored(tmp_path)[_URL]["experience_level"] == EXPERIENCE_UNREADABLE


def test_a_rejected_unreadable_row_costs_no_request(env: dict[str, Any]) -> None:
    """Layer 4 runs before Layer 5, so a row the owner rejected is never re-fetched."""
    site, tmp_path = env["site"], env["tmp_path"]
    _run(tmp_path)
    with JobStore(tmp_path / "jobs.sqlite3") as store:
        store.set_status([_URL], "rejected")
    _run(tmp_path)
    assert site.fetches[_URL] == 1


def test_a_deferred_job_against_a_shell_is_not_stored_rejected(env: dict[str, Any]) -> None:
    env["extracted"][0]["location"] = "Home based - Worldwide"
    site, tmp_path = env["site"], env["tmp_path"]
    summary = _run(tmp_path)
    assert _URL not in _stored(tmp_path), "unverified: dropped this run, never rejected"
    assert summary.unreadable_pages == (UnreadablePages(_SOURCE, 1, 1),)

    env["site"].page = posting("You will work from our Berlin office.")
    _run(tmp_path)
    assert site.fetches[_URL] == 2, "retried next run, because nothing was stored"
    assert _stored(tmp_path)[_URL]["status"] == "new"


def test_a_stored_new_job_that_a_recheck_reads_and_rejects_is_rejected(
    env: dict[str, Any],
) -> None:
    """A revived row is 'new'. Judged against a real page and found wanting, it
    must not stay in the review sheet next to a "senior" level."""
    site, tmp_path = env["site"], env["tmp_path"]
    _run(tmp_path)
    site.page = posting("We require 8+ years of experience in the field.")
    _run(tmp_path)
    job = _stored(tmp_path)[_URL]
    assert job["status"] == "rejected"
    assert job["experience_level"] == "senior (8+yr)"


def test_a_stored_seen_job_that_a_recheck_rejects_keeps_the_owners_status(
    env: dict[str, Any],
) -> None:
    site, tmp_path = env["site"], env["tmp_path"]
    _run(tmp_path)
    with JobStore(tmp_path / "jobs.sqlite3") as store:
        store.set_status([_URL], "shortlisted")
    site.page = posting("We require 8+ years of experience in the field.")
    _run(tmp_path)
    assert _stored(tmp_path)[_URL]["status"] == "shortlisted"


def test_a_stored_job_dropped_as_unverified_shows_as_unchecked_not_blank(
    env: dict[str, Any],
) -> None:
    """A revived row is stored with no level. If its deferred location cannot be
    verified it is dropped for the run and its row is untouched, so without this
    the sheet would show it blank."""
    tmp_path = env["tmp_path"]
    _run(tmp_path)  # stored, unchecked
    with JobStore(tmp_path / "jobs.sqlite3") as store:
        store._c().execute("UPDATE jobs SET experience_level = ''")  # as after a revival
    env["extracted"][0]["location"] = "Home based - Worldwide"
    _run(tmp_path)
    job = _stored(tmp_path)[_URL]
    assert (job["status"], job["experience_level"]) == ("new", EXPERIENCE_UNREADABLE)


def test_marking_unlabelled_never_overwrites_a_level(tmp_path: Path) -> None:
    with JobStore(tmp_path / "jobs.sqlite3") as store:
        run_id = store.begin_run()
        store.upsert_jobs(
            [
                {"dedupe_key": "a", "source_name": "s", "experience_level": "junior (<=2yr)"},
                {"dedupe_key": "b", "source_name": "s"},
            ],
            run_id,
        )
        assert store.mark_unlabelled(["a", "b"], EXPERIENCE_UNREADABLE) == 1
        levels = {j["dedupe_key"]: j["experience_level"] for j in store.all_jobs()}
    assert levels == {"a": "junior (<=2yr)", "b": EXPERIENCE_UNREADABLE}


# --- a description the reader supplied (SP4d) --------------------------------

# jobsinlund's shortest posting was about 420 characters (SP4b): a real posting,
# under the threshold a fetched page is held to.
_SUPPLIED = ("A short posting for an analyst role in our team. " * 9)[:420]


def test_a_supplied_description_is_exempt_from_the_length_test_only() -> None:
    assert len(_SUPPLIED) == 420 < MIN_READABLE_CHARS
    assert is_unreadable(_SUPPLIED)
    assert not is_unreadable(_SUPPLIED, supplied=True)
    assert is_unreadable("Please enable JavaScript to run this app.", supplied=True)
    assert is_unreadable("   ", supplied=True)


def test_layer5_reads_a_supplied_description_and_fetches_nothing() -> None:
    (kept,), excluded = apply_detail_filter([_job(description_text=_SUPPLIED)], _boom)
    assert not excluded
    assert kept["experience_level"] == "unspecified", "read, not failed and not unreadable"
    assert kept["description_text"] == _SUPPLIED
    assert kept["description_fetched_at"]


def test_a_supplied_description_is_judged_like_a_page() -> None:
    text = _SUPPLIED + " You bring 5+ years of experience."
    kept, (dropped,) = apply_detail_filter([_job(description_text=text)], _boom)
    assert not kept
    assert dropped["experience_level"] == "senior (5+yr)"


def test_an_empty_supplied_description_falls_back_to_the_page() -> None:
    (kept,), _ = _detail(_job(description_text="  "), posting("No requirement stated."))
    assert kept["experience_level"] == "unspecified"


def test_a_supplied_description_is_stored_and_not_fetched_on_the_next_run(
    env: dict[str, Any],
) -> None:
    """Before SP4d's pipeline change a stored description under the threshold
    was re-checked every run, and a short supplied one would never have settled.
    """
    site, tmp_path = env["site"], env["tmp_path"]
    env["extracted"][0]["description_text"] = _SUPPLIED

    first = _run(tmp_path)
    job = _stored(tmp_path)[_URL]
    assert (job["experience_level"], job["description_text"]) == ("unspecified", _SUPPLIED)
    assert first.unreadable_pages == ()

    second = _run(tmp_path)
    assert second.jobs_stored_rechecked == 0, "judged on this text already"
    assert site.fetches[_URL] == 0, "no run fetches a page the reader already supplied"


def test_a_stored_shell_is_rejudged_from_the_supplied_text_without_a_fetch(
    env: dict[str, Any],
) -> None:
    """jobsinlund's stored rows hold the shells of the pages fetched before SP4d."""
    site, tmp_path = env["site"], env["tmp_path"]
    _run(tmp_path)
    assert site.fetches[_URL] == 1

    env["extracted"][0]["description_text"] = _SUPPLIED
    second = _run(tmp_path)
    assert second.jobs_stored_rechecked == 1
    assert site.fetches[_URL] == 1
    assert _stored(tmp_path)[_URL]["description_text"] == _SUPPLIED


def test_a_skipped_stored_job_keeps_the_description_it_was_judged_on(
    env: dict[str, Any],
) -> None:
    """A rejected row skips Layer 5, so the reader's text must not replace its own."""
    tmp_path = env["tmp_path"]
    env["site"].page = posting("You bring 6+ years of experience.")
    _run(tmp_path)
    before = _stored(tmp_path)[_URL]
    assert before["status"] == "rejected"

    env["extracted"][0]["description_text"] = _SUPPLIED
    _run(tmp_path)
    assert _stored(tmp_path)[_URL]["description_text"] == before["description_text"]
